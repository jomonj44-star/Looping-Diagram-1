"""Reads the marked-up PDF and builds the circuit data (no GUI code here)."""
import math, re, collections
import pdfplumber


def _rc(c):
    if c is None:
        return None
    if isinstance(c, (int, float)):
        c = (c,)
    return tuple(round(float(x), 2) for x in c)


def _decode(b):
    if b is None:
        return ""
    if isinstance(b, bytes):
        if b[:2] in (b"\xfe\xff", b"\xff\xfe"):
            return b.decode("utf-16", "ignore").strip()
        return b.decode("latin1", "ignore").strip()
    return str(b).strip()


def to_view(page, x, y):
    """Annotation (raw PDF) coordinates -> pdfplumber view coordinates (x, top)."""
    r = page.rotation % 360
    mb = page.page_obj.mediabox
    w, h = float(mb[2] - mb[0]), float(mb[3] - mb[1])
    if r == 0:
        return (x, h - y)
    if r == 90:
        return (y, x)
    if r == 180:
        return (w - x, y)
    return (h - y, w - x)


def text_lines(page, gap):
    chars = sorted(page.chars, key=lambda c: (c["bottom"], c["x0"]))
    rows = []
    for c in chars:
        if rows and abs(rows[-1][0] - c["bottom"]) < 0.8:
            rows[-1][1].append(c)
        else:
            rows.append([c["bottom"], [c]])
    lines = []
    for _, cs in rows:
        cs.sort(key=lambda c: c["x0"])
        seg = [cs[0]]
        for c in cs[1:]:
            if c["x0"] - seg[-1]["x1"] > gap:
                lines.append(seg)
                seg = [c]
            else:
                seg.append(c)
        lines.append(seg)
    out = []
    for seg in lines:
        t = "".join(c["text"] for c in seg).strip()
        if t:
            out.append(dict(text=t, x0=min(c["x0"] for c in seg), x1=max(c["x1"] for c in seg),
                            top=min(c["top"] for c in seg), bottom=max(c["bottom"] for c in seg)))
    return out


def rect_dist(pt, r):
    dx = max(r[0] - pt[0], 0, pt[0] - r[2])
    dy = max(r[1] - pt[1], 0, pt[1] - r[3])
    return math.hypot(dx, dy)


def read_pdf(path, S, log=print):
    pdf = pdfplumber.open(path)
    page = pdf.pages[int(S["page"])]
    # 1. new tags (annotations)
    news = []
    for a in page.annots:
        d = a["data"]
        if S["annot_subtype"] not in str(d.get("Subtype")):
            continue
        txt = _decode(d.get("Contents"))
        if not txt:
            continue
        cl = d.get("CL")
        if cl and len(cl) >= 2:
            tip = to_view(page, float(cl[0]), float(cl[1]))
        else:
            tip = ((a["x0"] + a["x1"]) / 2, (a["top"] + a["bottom"]) / 2)
        news.append(dict(new=" ".join(txt.split()), tip=tip))
    log(f"{len(news)} new tags (annotations)")
    # 2. old tag boxes
    lines = text_lines(page, float(S["text_gap"]))
    rx = re.compile(S["old_tag_regex"])
    boxes = []
    for ln in lines:
        if not rx.match(ln["text"]):
            continue
        cand = [l for l in lines if l is not ln and 0 <= l["top"] - ln["bottom"] < 3
                and l["x1"] > ln["x0"] - 25 and l["x0"] < ln["x1"] + 25]
        sec = min(cand, key=lambda l: abs(l["x0"] - ln["x0"])) if cand else None
        x0 = min([ln["x0"]] + ([sec["x0"]] if sec else []))
        x1 = max([ln["x1"]] + ([sec["x1"]] if sec else []))
        bt = max(ln["bottom"], sec["bottom"] if sec else 0)
        boxes.append(dict(old=ln["text"], old2=sec["text"] if sec else "", rect=(x0, ln["top"], x1, bt)))
    log(f"{len(boxes)} old tag boxes")
    # 3. leader lines
    lc = _rc(S["leader_color"])
    segs = []
    for o in page.lines + page.curves:
        if _rc(o.get("stroking_color")) != lc:
            continue
        ps = [(float(p[0]), float(p[1])) for p in (o.get("pts") or [(o["x0"], o["top"]), (o["x1"], o["bottom"])])]
        for a, b in zip(ps, ps[1:]):
            segs.append((a, b))
    parent = list(range(len(segs)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    grid = collections.defaultdict(list)
    for i, (a, b) in enumerate(segs):
        for p in (a, b):
            grid[(round(p[0]), round(p[1]))].append((i, p))
    for i, (a, b) in enumerate(segs):
        for p in (a, b):
            cx, cy = round(p[0]), round(p[1])
            for dx in (-1, 0, 1):
                for dy in (-1, 0, 1):
                    for j, q in grid[(cx + dx, cy + dy)]:
                        if j != i and math.dist(p, q) < 0.6:
                            parent[find(i)] = find(j)
    comps = collections.defaultdict(list)
    for i in range(len(segs)):
        comps[find(i)].append(i)
    # 4. match each new tag to an old box
    poles, warns = [], []
    used = collections.Counter()
    for n in news:
        best = None
        for i, (a, b) in enumerate(segs):
            for p in (a, b):
                dd = math.dist(p, n["tip"])
                if best is None or dd < best[0]:
                    best = (dd, i, p)
        pole = dict(n, old=None, old2="", d=None, box=None)
        if best and best[0] < 20:
            _, i, p = best
            eps = [q for j in comps[find(i)] for q in segs[j]]
            far = max(eps, key=lambda q: math.dist(q, p))
            bi = min(range(len(boxes)), key=lambda k: rect_dist(far, boxes[k]["rect"])) if boxes else None
            if bi is not None:
                d = rect_dist(far, boxes[bi]["rect"])
                if d <= float(S["match_tol"]):
                    pole.update(old=boxes[bi]["old"], old2=boxes[bi]["old2"], d=d, box=bi)
                    used[bi] += 1
        if pole["old"] is None:
            warns.append(f"No old tag found for {n['new']}")
        poles.append(pole)
    for bi, c in used.items():
        if c > 1:
            warns.append(f"Old tag box '{boxes[bi]['old']}' is used by {c} new tags")
    log(f"matched {sum(1 for p in poles if p['old'])} of {len(poles)}")
    return dict(poles=poles, warnings=warns, page_size=(page.width, page.height))


def detect_type(S, pole):
    txt = (pole.get("old2") or "") + " " + (pole.get("old") or "")
    for ft in sorted(S["fixture_types"], key=lambda f: -len(f["match"])):
        if ft["match"] and ft["match"] in txt:
            return ft["name"]
    return None


def build(S, data):
    unit = 25.4 / 72.0 * float(S["scale"]) / 1000.0 * float(S["paper_factor"]) * float(S["length_multiplier"])
    rx = re.compile(S["tag_regex"])
    ph = {p: i for i, p in enumerate(S["phase_order"])}
    nph = len(ph)
    checks = list(data["warnings"])
    groups = collections.defaultdict(list)
    for p in data["poles"]:
        m = rx.match(p["new"])
        if not m:
            checks.append(f"Tag format not recognised: {p['new']}")
            continue
        g = m.groupdict()
        fx = [(int(g["no"]), g["ph"])]
        if g.get("no2"):
            fx.append((int(g["no2"]), g["ph2"]))
            if g.get("cct2") and g["cct2"] != g["cct"]:
                checks.append(f"Circuit mismatch inside tag: {p['new']}")
        if any(f[1] not in ph for f in fx):
            checks.append(f"Phase not in phase order: {p['new']}")
            continue
        p = dict(p, panel=g["panel"], cct=g["cct"], fx=fx, ld=detect_type(S, p))
        if p["ld"] is None:
            checks.append(f"Fixture type not found for {p['new']} (old tag text: {p['old2']})")
        groups[(g["panel"], g["cct"])].append(p)
    pos = lambda f: (f[0] - 1) * nph + ph[f[1]]
    circuits = {}
    for k in sorted(groups):
        v = sorted(groups[k], key=lambda p: pos(p["fx"][0]))
        segs = [(math.dist(a["tip"], b["tip"]) * unit + float(S["extra_per_span_m"])) for a, b in zip(v, v[1:])]
        cnt = collections.Counter(pos(f) for p in v for f in p["fx"])
        dups = [f"L{S['phase_order'][x % nph]}-{x // nph + 1:02d}" for x, c in sorted(cnt.items()) if c > 1]
        gaps = [f"L{S['phase_order'][x % nph]}-{x // nph + 1:02d}" for x in range(0, max(cnt) + 1) if x not in cnt]
        if dups:
            checks.append(f"{k[0]}/{k[1]}: duplicate fixture {', '.join(dups)}")
        if gaps:
            checks.append(f"{k[0]}/{k[1]}: missing in sequence {', '.join(gaps)}")
        for i, L in enumerate(segs):
            if L > float(S["long_span_warn_m"]):
                checks.append(f"{k[0]}/{k[1]}: long span {L:.1f} m between pole {i + 1} and {i + 2}")
        size = S["cable_sizes"].get(f"{k[0]}/{k[1]}", S["cable_sizes"].get("default", ""))
        types = collections.OrderedDict((ft["name"], [0, 0]) for ft in S["fixture_types"])
        for p in v:
            t = p["ld"] or "?"
            types.setdefault(t, [0, 0])
            types[t][0] += 1
            types[t][1] += len(p["fx"])
        circuits[k] = dict(key=k, poles=v, segs=segs, length=sum(segs), cable=size, types=types,
                           fixtures=sum(len(p["fx"]) for p in v))
    return dict(circuits=circuits, checks=checks, unit=unit)


def overall(S, built):
    types = collections.OrderedDict((ft["name"], [0, 0]) for ft in S["fixture_types"])
    cable = collections.OrderedDict()
    for c in built["circuits"].values():
        for t, (a, b) in c["types"].items():
            types.setdefault(t, [0, 0])
            types[t][0] += a
            types[t][1] += b
        cable.setdefault(c["cable"], [0, 0.0])
        cable[c["cable"]][0] += 1
        cable[c["cable"]][1] += c["length"]
    types = collections.OrderedDict((t, v) for t, v in types.items() if v[0])
    return types, cable


def cable_tag(S, c, n):
    return S["cable_tag"].format(panel=c["key"][0], cct=c["key"][1], n=n)


def export_xlsx(path, S, built):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = "MTO by circuit"
    types, cable = overall(S, built)
    names = list(types)
    ws.append(["Panel", "Circuit", "Poles", "Fixtures"] + [f"{t} (poles/fixtures)" for t in names] +
              ["Lighting JB", "Cable size", "Length (m)"])
    for c in built["circuits"].values():
        ws.append([c["key"][0], c["key"][1], len(c["poles"]), c["fixtures"]] +
                  [f"{c['types'].get(t, [0, 0])[0]}/{c['types'].get(t, [0, 0])[1]}" for t in names] +
                  [len(c["poles"]), c["cable"], round(c["length"], 1)])
    ws.append(["TOTAL", "", sum(len(c["poles"]) for c in built["circuits"].values()),
               sum(c["fixtures"] for c in built["circuits"].values())] +
              [f"{types[t][0]}/{types[t][1]}" for t in names] +
              [sum(len(c["poles"]) for c in built["circuits"].values()), "",
               round(sum(c["length"] for c in built["circuits"].values()), 1)])
    ws2 = wb.create_sheet("Cable by size")
    ws2.append(["Cable size", "Circuits", "Total length (m)"])
    for k, (n, L) in cable.items():
        ws2.append([k, n, round(L, 1)])
    ws3 = wb.create_sheet("Spans")
    ws3.append(["Panel", "Circuit", "Cable tag", "From JB tag", "To JB tag", "Cable size", "Length (m)"])
    for c in built["circuits"].values():
        prev = None
        for i, p in enumerate(c["poles"]):
            jb = p["new"] + S["jb_suffix"]
            ws3.append([c["key"][0], c["key"][1], cable_tag(S, c, i + 1), prev or c["key"][0] + " (panel)", jb,
                        c["cable"], round(c["segs"][i - 1], 1) if i else None])
            prev = jb
    ws4 = wb.create_sheet("Tag list")
    ws4.append(["Old tag", "Old tag line 2", "New tag", "JB tag", "Fixture type"])
    for c in built["circuits"].values():
        for p in c["poles"]:
            ws4.append([p["old"], p["old2"], p["new"], p["new"] + S["jb_suffix"], p["ld"]])
    for w in wb.worksheets:
        for cell in w[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = max(12, min(40, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(path)


def extract_symbol(page, cx, cy, S, name="NEW"):
    """Copy the vector drawing around view position (cx, cy) as a symbol template."""
    P = S["pick"]
    R, MX = float(P["radius"]), float(P["max_size"])
    ign = [_rc(x) for x in P["ignore_colors"]]
    jbc = _rc(P["jb_color"])
    keep = []
    for o in page.lines + page.curves + page.rects:
        if o["x0"] < cx - R or o["x1"] > cx + R or o["top"] < cy - R or o["bottom"] > cy + R:
            continue
        if o["x1"] - o["x0"] > MX or o["bottom"] - o["top"] > MX:
            continue
        s, f = _rc(o.get("stroking_color")), _rc(o.get("non_stroking_color"))
        if s in ign or f in ign:
            continue
        if s == jbc and o["object_type"] == "line" and math.hypot(o["x1"] - o["x0"], o["bottom"] - o["top"]) > 4.6:
            continue
        keep.append(o)
    if not keep:
        return None

    def hexc(c):
        if c is None:
            return None
        if len(c) == 1:
            return "k" if c[0] == 0.0 else "#808080"
        if len(c) == 4:
            return "k"
        return "#%02x%02x%02x" % tuple(int(v * 255) for v in c[:3])
    prims, xs, ys = [], [], []
    for o in keep:
        pts = [(float(a), float(b)) for a, b in (o.get("pts") or [(o["x0"], o["top"]), (o["x1"], o["bottom"])])]
        fill = hexc(_rc(o.get("non_stroking_color"))) if o.get("fill") and len(pts) > 2 else None
        stroke = hexc(_rc(o.get("stroking_color"))) if o.get("stroke", True) else None
        prims.append(dict(pts=pts, stroke=stroke, fill=fill))
        xs += [q[0] for q in pts]
        ys += [q[1] for q in pts]
    mx, my = (min(xs) + max(xs)) / 2, (min(ys) + max(ys)) / 2
    for p in prims:
        p["pts"] = [(a - mx, b - my) for a, b in p["pts"]]
    # JB square = group of JB-coloured primitives (largest square-like group)
    jbh = hexc(jbc)
    pur = [p for p in prims if p["stroke"] == jbh]
    groups = []
    for p in pur:
        bx = (min(q[0] for q in p["pts"]), min(q[1] for q in p["pts"]), max(q[0] for q in p["pts"]), max(q[1] for q in p["pts"]))
        for g in groups:
            if bx[0] <= g[2] + 0.3 and bx[2] >= g[0] - 0.3 and bx[1] <= g[3] + 0.3 and bx[3] >= g[1] - 0.3:
                g[0], g[1], g[2], g[3] = min(g[0], bx[0]), min(g[1], bx[1]), max(g[2], bx[2]), max(g[3], bx[3])
                break
        else:
            groups.append(list(bx))
    groups = [g for g in groups if 1.0 < min(g[2] - g[0], g[3] - g[1]) < 4.0 and max(g[2] - g[0], g[3] - g[1]) <= 6.0]
    jb = [0.0, 0.0, 2.0]
    if groups:
        g = max(groups, key=lambda g: (round(min(g[2] - g[0], g[3] - g[1]), 1), -abs((g[1] + g[3]) / 2)))
        w_, h_ = g[2] - g[0], g[3] - g[1]
        sd = min(w_, h_)
        # two squares stacked/side by side: the first (top / left) one is the JB
        jb = [(g[0] + g[2]) / 2 if w_ <= h_ else g[0] + sd / 2, g[1] + sd / 2 if h_ >= w_ else (g[1] + g[3]) / 2, sd]
    return dict(name=name, w=max(xs) - min(xs), h=max(ys) - min(ys), jb=jb, prims=prims)
