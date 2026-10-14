"""Reads the 'Street Lighting Circuiting - Load Balance' workbook into circuit blocks (poles, parents, lengths)."""
import re
import openpyxl


def _s(v):
    return "" if v is None else str(v).strip()


def norm_tag(t):
    t = re.sub(r"\s*/\s*", "/", _s(t))
    return re.sub(r"\s+", "", t)


def _area(ws, col):
    for c in range(col, col + 1):
        v = ws.cell(1, c).value
        if v:
            return " | ".join(x.strip() for x in str(v).splitlines() if x.strip())
    return ""


def read_workbook(path, log=print):
    wb = openpyxl.load_workbook(path, data_only=True)
    blocks, warns = [], []
    for ws in wb.worksheets:
        # find table headers
        heads = []
        for r in range(1, ws.max_row + 1):
            for c in range(1, ws.max_column + 1):
                h0 = _s(ws.cell(r, c).value).lower().replace(".", "").replace(" ", "")
                if h0 == "slno" and "panel" in _s(ws.cell(r, c + 1).value).lower() and \
                        "fixture" in _s(ws.cell(r, c + 3).value).lower():
                    heads.append((r, c))
                elif h0.startswith("panelnumber") and "circuit" in _s(ws.cell(r, c + 1).value).lower() and \
                        "fixture" in _s(ws.cell(r, c + 2).value).lower() and \
                        (c == 1 or _s(ws.cell(r, c - 1).value).lower().replace(".", "").replace(" ", "") != "slno"):
                    # header row without the SL.NO label (column titles shifted by one); the data below keeps the normal layout
                    heads.append((r, c))
        for (hr, c0) in heads:
            title = ""
            for rr in range(hr - 1, max(hr - 4, 0), -1):
                t = _s(ws.cell(rr, c0).value)
                if "CIRCUITING" in t.upper():
                    title = t
                    break
            if not title:   # later tables on the sheet have no title row: use the nearest earlier one in this column
                for rr in range(hr - 1, 0, -1):
                    t = _s(ws.cell(rr, c0).value)
                    if "CIRCUITING" in t.upper():
                        title = t
                        break
            rows, total = [], {}
            r = hr + 1
            blank = 0
            while r <= ws.max_row:
                v = [ws.cell(r, c0 + i).value for i in range(10)]
                if _s(v[2]).upper() == "TOTAL LOAD":
                    # total load rows: C01L1.. watts; length total sits in the LENGTH column on one of the rows
                    for k in range(4):
                        rr = ws.cell(r + k, c0 + 3).value or ws.cell(r + k, c0 + 2).value
                        w = ws.cell(r + k, c0 + 5).value
                        if rr and isinstance(w, (int, float)) and _s(rr).upper() != "TOTAL LOAD":
                            total[_s(rr)] = w
                    break
                if _s(v[0]).lower().replace(".", "").replace(" ", "") == "slno":
                    break
                if _s(v[3]):
                    rows.append(dict(tag=norm_tag(v[3]), panel=_s(v[1]).replace(" ", ""), circuit=_s(v[2]), frm=norm_tag(v[4]),
                                     watt=v[5] or 0, seq=_s(v[6]), phase=_s(v[7]), type=_s(v[8]),
                                     length=v[9] if isinstance(v[9], (int, float)) else None, row=r))
                    blank = 0
                else:
                    blank += 1
                    if blank > 3 and rows:
                        break
                r += 1
            if not rows:
                continue
            # length total printed under the table (first numeric in LENGTH col after the last data row)
            xl_total = None
            for rr in range(rows[-1]["row"] + 1, min(rows[-1]["row"] + 8, ws.max_row + 1)):
                x = ws.cell(rr, c0 + 9).value
                if isinstance(x, (int, float)):
                    xl_total = x
                    break
            blocks.append(dict(sheet=ws.title, area=_area(ws, c0), title=title, rows=rows, loads=total, xl_length=xl_total, col=c0))
    return blocks


def build_poles(block):
    """Poles = rows with a fixture type; following rows without type are further fixtures on the same pole."""
    rows = block["rows"]
    poles, warns = [], []
    for rw in rows:
        if rw["type"] or not poles:
            poles.append(dict(rows=[rw], type=rw["type"] or None, frm=rw["frm"], length=rw["length"]))
        else:
            poles[-1]["rows"].append(rw)
    owner = {}
    for i, p in enumerate(poles):
        names = [r["tag"] for r in p["rows"]]
        first = names[0]
        p["fx"] = names
        p["tag"] = first if len(names) == 1 else first + "".join("/" + n.split("/")[-1] for n in names[1:])
        p["watt"] = sum(r["watt"] or 0 for r in p["rows"])
        p["idx"] = i
        for n in names:
            owner[n] = i
    panel = block["rows"][0]["panel"]
    for i, p in enumerate(poles):
        f = p["frm"]
        if not f or "/" not in f:           # panel (or empty on the first pole)
            p["parent"] = -1 if (i == 0 or f) else i - 1
            if not f and i > 0:
                warns.append(f"{p['tag']}: FROM is empty, assumed previous pole")
        elif f in owner and owner[f] < i:
            p["parent"] = owner[f]
        else:
            p["parent"] = i - 1
            warns.append(f"{p['tag']}: FROM '{f}' not found before this pole, assumed previous pole")
    return poles, warns


TAGRX = re.compile(r"^(?P<panel>[A-Za-z0-9]+)/(?P<cct>\d+)L(?P<ph>\d)-(?P<no>\d+)")


def build_excel(S, blocks, log=print):
    """Turns workbook blocks into the same 'built' structure used by the renderer."""
    import collections
    circuits = collections.OrderedDict()
    checks = []
    for bi, b in enumerate(blocks):
        poles, w = build_poles(b)
        checks += [f"{b['sheet']}: {x}" for x in w]
        first = b["rows"][0]
        m = TAGRX.match(first["tag"])
        panel = m.group("panel") if m else first["panel"]
        cct = m.group("cct") if m else re.sub(r"\D", "", first["circuit"])
        kind = "EMERGENCY" if "EMERGENCY" in b["title"].upper() or "BNL" in panel else "NORMAL"
        label = f"{panel} / {cct}   ({kind.title()}, {len(poles)} poles)"
        types = collections.OrderedDict((ft["name"], [0, 0]) for ft in S["fixture_types"])
        known = {ft["name"] for ft in S["fixture_types"]}
        for p in poles:
            t = p["type"] if p["type"] in known else (p["type"] or "?")
            p["ld"] = t if t in known else None
            p["new"] = p["tag"]
            p["old"] = None
            types.setdefault(t, [0, 0])
            types[t][0] += 1
            types[t][1] += len(p["rows"])
            if p["type"] and p["type"] not in known:
                checks.append(f"{p['tag']}: fixture type '{p['type']}' is not in Settings -> Fixture types")
            if not p["type"]:
                checks.append(f"{p['tag']}: no fixture type in Excel")
            if p["parent"] >= 0 and p["length"] is None:
                checks.append(f"{p['tag']}: no length in Excel")
        length = sum(p["length"] or 0 for p in poles)
        phase = collections.OrderedDict([("L1", 0), ("L2", 0), ("L3", 0)])
        for r in b["rows"]:
            mm = TAGRX.match(r["tag"])
            ph = "L" + mm.group("ph") if mm else (r["phase"] or "?")
            if r["phase"] and r["phase"] != ph:
                checks.append(f"{r['tag']}: PHASES column says {r['phase']} but the tag says {ph} (tag used)")
            phase[ph] = phase.get(ph, 0) + (r["watt"] or 0)
        if b["xl_length"] is not None and abs(b["xl_length"] - length) > 0.5:
            checks.append(f"{panel}/{cct}: sum of lengths {length} differs from Excel total {b['xl_length']}")
        tags = [r["tag"] for r in b["rows"]]
        dups = [t for t, n in collections.Counter(tags).items() if n > 1]
        if dups:
            checks.append(f"{panel}/{cct}: duplicate fixture tag {', '.join(dups)}")
        size = S["cable_sizes"].get(f"{panel}/{cct}", S["cable_sizes"].get("default", ""))
        circuits[bi + 1] = dict(key=(panel, cct), label=label, kind=kind, sheet=b["sheet"], area=b["area"], title=b["title"],
                                poles=poles, types=types, fixtures=sum(len(p["rows"]) for p in poles),
                                length=length, segs=[p["length"] or 0 for p in poles], cable=size, phase=phase,
                                excel_loads=b["loads"])
    return dict(circuits=circuits, checks=checks, source="excel")


def export_xlsx(path, S, built):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    import core
    wb = Workbook()
    ws = wb.active
    ws.title = "MTO by circuit"
    types, cable = core.overall(S, built)
    names = list(types)
    ws.append(["Panel", "Circuit", "Type", "Area", "Poles", "Fixtures"] + [f"{t} (poles/fixtures)" for t in names] +
              ["Lighting JB", "L1 (W)", "L2 (W)", "L3 (W)", "Load (W)", "Cable size", "Cable length (m)"])
    for c in built["circuits"].values():
        ph = c["phase"]
        ws.append([c["key"][0], c["key"][1], c["kind"].title(), c["area"].split("|")[0].split(":", 1)[-1].strip(), len(c["poles"]), c["fixtures"]] +
                  [f"{c['types'].get(t, [0, 0])[0]}/{c['types'].get(t, [0, 0])[1]}" for t in names] +
                  [len(c["poles"]), ph.get("L1", 0), ph.get("L2", 0), ph.get("L3", 0), sum(ph.values()), c["cable"], c["length"]])
    ws2 = wb.create_sheet("Cable runs")
    ws2.append(["Panel", "Circuit", "Cable tag", "From", "To (JB tag)", "Fixture type", "Cable size", "Length (m)"])
    import render_tree
    for c in built["circuits"].values():
        order = render_tree.layout(c["poles"])
        for n, i in enumerate(order, 1):
            p = c["poles"][i]
            par = p["parent"]
            frm = (c["key"][0] + " (panel)") if par < 0 else c["poles"][par]["tag"] + S["jb_suffix"]
            ws2.append([c["key"][0], c["key"][1], core.cable_tag(S, c, n), frm, p["tag"] + S["jb_suffix"], p["type"], c["cable"], p["length"]])
    ws3 = wb.create_sheet("Poles")
    ws3.append(["Panel", "Circuit", "Pole tag", "JB tag", "Fixture tags", "Fixture type", "Load (W)", "Length (m)", "From (as in Excel)"])
    for c in built["circuits"].values():
        for p in c["poles"]:
            ws3.append([c["key"][0], c["key"][1], p["tag"], p["tag"] + S["jb_suffix"], ", ".join(p["fx"]), p["type"], p["watt"], p["length"], p["frm"]])
    ws4 = wb.create_sheet("Checks")
    ws4.append(["Check"])
    for x in built["checks"]:
        ws4.append([x])
    for w in wb.worksheets:
        for cell in w[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = max(12, min(60, max(len(str(c.value or "")) for c in col) + 2))
    wb.save(path)
